"""New-tool tests (SQL functions and generated MCP servers) - no network.

Run: python test_forge.py

Each "must be refused" case below is something a model-written tool has no business
doing; each "must pass" case is something a real tool needs. The last case builds
the files for a good tool and loads them with the real FastMCP library when it is
available (set FORGE_PYTHON to a Python that has fastmcp, openpyxl and httpx; it is
skipped, and says so, otherwise).
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile

import forge
from dbx import DbxError

CASES = []


def case(name):
    def deco(fn):
        CASES.append((name, fn))
        return fn

    return deco


GOOD = '''
import io

import openpyxl

COLUMNS = ("position", "rate")


def read_rate_card(file_path: str, sheet: str = "") -> dict:
    """Read a rate card spreadsheet. file_path: the file in the rate card folder. sheet: optional sheet name."""
    wb = openpyxl.load_workbook(io.BytesIO(volume_read(file_path)), data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return {"rows": rows[:200], "count": len(rows)}


def save_plan(plan: dict, file_path: str) -> dict:
    """Save a plan as a small spreadsheet. plan: rows to save. file_path: where to put it."""
    wb = openpyxl.Workbook()
    for row in plan.get("rows", []):
        wb.active.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return {"path": volume_write(file_path, buf.getvalue())}
'''


def tool(code=GOOD, **kw):
    base = {
        "kind": "mcp", "slug": "rate-card-tool", "name": "Rate card tool", "description": "Reads and saves rate card files.",
        "abilities": [
            {"name": "read_rate_card", "description": "Reads a rate card.", "changes_data": False},
            {"name": "save_plan", "description": "Saves a plan as a spreadsheet.", "changes_data": True}],
        "code": code, "hosts": [], "secrets": [],
        "volumes": [{"volume": "main.sales.rates", "access": "read"}, {"volume": "main.sales.results", "access": "write"}],
    }
    base.update(kw)
    return forge.clean_item(base)


def says(item, fragment):
    assert any(fragment in p for p in item["problems"]), (fragment, item["problems"])


@case("a tool that reads and saves spreadsheets passes, and says what it can reach")
def _():
    t = tool()
    assert t["problems"] == [], t["problems"]
    r = t["report"]
    assert r["reads_files"] and r["saves_files"] and not r["calls_internet"]
    assert r["packages"] == ["openpyxl"]  # from its imports, not from anything the model wrote
    assert t["fingerprint"] and len(t["fingerprint"]) == 64


@case("things a tool has no business doing are refused")
def _():
    cases = {
        "eval": ("def read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return eval(file_path)\n", "eval"),
        "open": ("def read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {'t': open(file_path).read()}\n", "open"),
        "getattr": ("def read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {'t': getattr(file_path, 'upper')()}\n", "getattr"),
        "subprocess": ("import subprocess\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n", "subprocess cannot be used"),
        "from os": ("from os import system\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n", "from os"),
        "os.system": ("import os\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    os.system('ls')\n    return {}\n", ".system is not allowed"),
        "dunder": ("def read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {'t': ().__class__.__bases__}\n", "__class__"),
        "import time work": ("import io\nBUF = io.BytesIO()\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n", "top level"),
        "script": ("def read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n\nif __name__ == '__main__':\n    pass\n", "top level"),
        "unlisted import": ("import requests\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n", "requests cannot be used"),
        "urllib": ("import urllib.request\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n", "urllib.request cannot be used"),
        "async": ("async def read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n", "ordinary functions"),
        "untyped": ("def read_rate_card(file_path):\n    \"\"\"x\"\"\"\n    return {}\n", "give a type"),
        "no docstring": ("def read_rate_card(file_path: str) -> dict:\n    return {}\n", "docstring"),
        "syntax": ("def read_rate_card(:\n", "mistake on line"),
    }
    for label, (code, fragment) in cases.items():
        t = tool(code=code, abilities=[{"name": "read_rate_card", "description": "Reads.", "changes_data": False}], volumes=[])
        try:
            says(t, fragment)
        except AssertionError as exc:
            raise AssertionError("%s: %s" % (label, exc))


@case("column_letter is refused (it fails on merged cells), and the safe way of getting a column letter passes")
def _():
    ab = [{"name": "read_rate_card", "description": "Reads.", "changes_data": False}]
    bad = ("import io\nimport openpyxl\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n"
           "    ws = openpyxl.load_workbook(io.BytesIO(volume_read(file_path))).active\n"
           "    return {'cols': [c.column_letter for c in ws[1]]}\n")
    t = tool(code=bad, abilities=ab, volumes=[{"volume": "main.sales.rates", "access": "read"}])
    says(t, "fails on merged cells")
    good = bad.replace("import openpyxl\n", "import openpyxl\nfrom openpyxl.utils import get_column_letter\n").replace("c.column_letter", "get_column_letter(c.column)")
    assert tool(code=good, abilities=ab, volumes=[{"volume": "main.sales.rates", "access": "read"}])["problems"] == []


@case("settings, sites and folders must be declared; anything else is refused")
def _():
    body = "import os\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {'k': %s}\n"
    ab = [{"name": "read_rate_card", "description": "Reads.", "changes_data": False}]
    t = tool(code=body % "os.environ['GAM_KEY']", abilities=ab, volumes=[], secrets=["gam-key"])
    assert t["problems"] == [], t["problems"]
    says(tool(code=body % "os.environ['GAM_KEY']", abilities=ab, volumes=[], secrets=[]), "GAM_KEY was not declared")
    says(tool(code=body % "dict(os.environ)", abilities=ab, volumes=[], secrets=["gam-key"]), "cannot be checked")
    says(tool(code=body % "os.environ[file_path]", abilities=ab, volumes=[], secrets=["gam-key"]), "cannot be checked")
    site = "import httpx\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return httpx.get('https://api.example.com/x').json()\n"
    says(tool(code=site, abilities=ab, volumes=[], hosts=[]), "api.example.com")
    assert tool(code=site, abilities=ab, volumes=[], hosts=["api.example.com"])["problems"] == []
    vol = "def read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {'b': volume_read('/Volumes/main/other/secret/x.xlsx')}\n"
    says(tool(code=vol, abilities=ab, volumes=[{"volume": "main.sales.rates", "access": "read"}]), "not in a folder this tool declared")
    says(tool(code=vol.replace("other/secret", "sales/rates/../../x"), abilities=ab, volumes=[{"volume": "main.sales.rates", "access": "read"}]), "not a folder path")


@case("saving files or changing things outside must be declared as changing data")
def _():
    save = "def save_it(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {'p': volume_write(file_path, b'x')}\n"
    ab = lambda changes: [{"name": "save_it", "description": "Saves.", "changes_data": changes}]  # noqa: E731
    says(tool(code=save, abilities=ab(False), volumes=[{"volume": "main.a.b", "access": "write"}]), "marked as changing data")
    says(tool(code=save, abilities=ab(True), volumes=[{"volume": "main.a.b", "access": "read"}]), "write access")
    assert tool(code=save, abilities=ab(True), volumes=[{"volume": "main.a.b", "access": "write"}])["problems"] == []
    post = "import httpx\n\ndef send_it(text: str) -> dict:\n    \"\"\"x\"\"\"\n    return httpx.post('https://hooks.example.com/x', json={'t': text}).json()\n"
    t = tool(code=post, abilities=[{"name": "send_it", "description": "Sends.", "changes_data": False}], volumes=[], hosts=["hooks.example.com"])
    says(t, "marked as changing data")
    assert t["report"]["may_change_outside"]


@case("the abilities and names have to line up with the code")
def _():
    says(tool(abilities=[{"name": "missing_one", "description": "x", "changes_data": False}], volumes=[]), "no function called missing_one")
    says(tool(abilities=[{"name": "Bad Name", "description": "x", "changes_data": False}]), "lowercase name")
    says(tool(slug="Bad Slug"), "lowercase letters")
    assert tool(code="")["problems"][-1] == "There is no code yet."


@case("a proposed tool's fingerprint and checks are recomputed, never taken from the browser")
def _():
    t = tool()
    forged = {**t, "problems": [], "fingerprint": "0" * 64, "report": {"kind": "mcp", "fake": True}}
    again = forge.clean_item({**forged, "code": GOOD + "\n\ndef extra(x: int) -> int:\n    \"\"\"x\"\"\"\n    return x\n"})
    assert again["fingerprint"] != t["fingerprint"] and "fake" not in again["report"]
    assert forge.clean_item(forged)["fingerprint"] == t["fingerprint"]
    assert forge.clean_item({"kind": "shell", "code": "rm -rf"}) is None and forge.clean_item("x") is None


@case("SQL functions: one plain CREATE FUNCTION, nothing that changes data")
def _():
    def fn(sql, **kw):
        base = {"kind": "uc_function", "name": "main.tools.fx", "description": "Converts money.", "sql": sql,
                "example": "SELECT main.tools.fx(10, 1.2)"}
        base.update(kw)
        return forge.clean_item(base)

    good = "CREATE FUNCTION main.tools.fx(amount DOUBLE, rate DOUBLE) RETURNS DOUBLE COMMENT 'Convert' RETURN amount * rate"
    assert fn(good)["problems"] == []
    assert fn(good + ";")["problems"] == []  # a trailing semicolon is fine
    says(fn(good.replace("CREATE FUNCTION", "CREATE OR REPLACE FUNCTION")), "not OR REPLACE")
    says(fn(good + "; DROP TABLE main.a.b"), "single statement")
    says(fn("CREATE FUNCTION main.tools.fx() RETURNS INT RETURN (SELECT 1) -- hi"), "no comments")
    says(fn("CREATE FUNCTION main.tools.fx() RETURNS INT LANGUAGE PYTHON AS $$ return 1 $$"), "LANGUAGE PYTHON")
    says(fn("CREATE FUNCTION main.tools.fx() RETURNS INT RETURN (DELETE FROM main.a.b)"), "DELETE")
    says(fn(good.replace("main.tools.fx", "main.tools.other")), "same as the function's name")
    says(fn("CREATE FUNCTION main.tools.fx() RETURNS INT"), "RETURN")
    says(fn(good, example="SELECT 1; DROP TABLE x"), "one plain SELECT")
    says(fn(good, example="SELECT 1"), "call the function")
    says(fn(good, name="bad name"), "catalog.schema.name")
    assert forge.fingerprint(fn(good)) != forge.fingerprint(fn(good + " + 1"))


@case("creating a function runs the SQL as given, then tries the example; a bad example is not fatal")
def _():
    sent = []

    def call(method, path, tok, **kw):
        sent.append((method, path, tok, (kw.get("json") or {}).get("statement")))
        stmt = (kw.get("json") or {}).get("statement", "")
        if stmt.upper().startswith("SELECT") and "boom" in stmt:
            return {"status": {"state": "FAILED", "error": {"message": "boom"}}}
        return {"status": {"state": "SUCCEEDED"}, "result": {"data_array": [["12.0"]]}}

    forge.call = call
    item = {"kind": "uc_function", "name": "main.tools.fx", "sql": "CREATE FUNCTION main.tools.fx() RETURNS INT RETURN 1;",
            "example": "SELECT main.tools.fx()"}
    out = forge.create_function(item, "wh123456", "ADMIN")
    assert out["created"] and out["example_result"] == [["12.0"]]
    assert sent[0][3] == "CREATE FUNCTION main.tools.fx() RETURNS INT RETURN 1" and all(s[2] == "ADMIN" for s in sent)
    out = forge.create_function({**item, "example": "SELECT boom"}, "wh123456", "ADMIN")
    assert out["created"] and out["example_error"] == "boom" and out["example_result"] is None
    try:
        forge.create_function(item, "", "ADMIN")
    except DbxError as exc:
        assert exc.status == 400
    else:
        raise AssertionError("a warehouse is required")

    def failing(method, path, tok, **kw):
        return {"status": {"state": "FAILED", "error": {"message": "no CREATE FUNCTION privilege"}}}

    forge.call = failing
    try:
        forge.create_function(item, "wh123456", "ADMIN")
    except DbxError as exc:
        assert "no CREATE FUNCTION privilege" in str(exc)  # Databricks' own words
    else:
        raise AssertionError("must raise")


@case("the files that get deployed: pinned packages from imports, settings as app resources, folders declared")
def _():
    t = tool(code=GOOD, secrets=[])
    files = forge.assemble(t, "a@x.io", "2026-10-08")
    assert set(files) == {"server.py", "requirements.txt", "app.yaml", "mcp.yaml", "README.md"}
    reqs = files["requirements.txt"].decode().split()
    assert reqs == ["fastmcp>=2.3", "httpx>=0.27", "openpyxl"], reqs
    server = files["server.py"].decode()
    assert "_READ_FOLDERS = ['main.sales.rates']" in server and "_WRITE_FOLDERS = ['main.sales.results']" in server
    assert "approved by a@x.io on 2026-10-08" in server
    assert "'readOnlyHint': True" in server and "'readOnlyHint': False" in server
    s = tool(code="import os\n\ndef read_rate_card(file_path: str) -> dict:\n    \"\"\"x\"\"\"\n    return {'k': os.environ['GAM_KEY']}\n",
             abilities=[{"name": "read_rate_card", "description": "Reads.", "changes_data": False}], volumes=[], secrets=["gam-key"])
    import yaml
    app = yaml.safe_load(forge.assemble(s, "a@x.io", "2026-10-08")["app.yaml"])
    assert app["env"] == [{"name": "GAM_KEY", "valueFrom": "gam-key"}] and app["command"] == ["python", "server.py"]
    card = yaml.safe_load(forge.assemble(s, "a@x.io", "2026-10-08")["mcp.yaml"])
    assert card["needs"]["secrets"] == ["gam-key"] and card["tools"][0]["changes_data"] is False
    entry = forge.entry_for(s)
    assert entry["app_name"] == "mcp-rate-card-tool" and entry["description"].startswith(forge.GEN_MARK)


@case("a quote in a name cannot break out of the generated file's docstring into code")
def _():
    import ast

    t = tool(name='Evil """ name', description='x """' + chr(10) + 'import os; os.system("x")' + chr(10) + '"""')
    server = forge.assemble(t, 'a"""@x.io', "2026-10-08")["server.py"].decode()
    tree = ast.parse(server)  # still valid, and the injected text stayed inside the docstring
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
             and isinstance(n.func.value, ast.Name) and n.func.value.id == "os" and n.func.attr == "system"]
    assert not calls


@case("it never overwrites: a name in use is refused unless it is this designer's own earlier attempt")
def _():
    calls = {}

    def call(method, path, tok, **kw):
        if path.startswith("/api/2.0/apps/"):
            if calls.get("app") == "missing":
                raise DbxError("nope", 404)
            return {"description": calls["app"]}
        if "/functions/" in path:
            if calls.get("fn") == "missing":
                raise DbxError("nope", 404)
            return {}
        return {}

    saved = (forge.call, forge.mcps.act, forge.mcps.catalog, forge.mcps.missing_secrets)
    try:
        forge.call = call
        forge.mcps.act = lambda method, path, tok, **kw: (call(method, path, tok, **kw), "you")
        forge.mcps.catalog = lambda refresh=False: ([{"slug": "gam-media-planner"}], "")
        forge.mcps.missing_secrets = lambda entry, tok="": []
        t = tool()
        calls["app"] = "somebody else's app"
        assert "was not made by this designer" in forge.name_problems(t, "T")[0]
        calls["app"] = forge.GEN_MARK + " Reads rate cards."
        assert forge.name_problems(t, "T") == []  # a retry of our own
        calls["app"] = "missing"
        assert forge.name_problems(t, "T") == []
        assert "already a ready-made tool" in forge.name_problems(tool(slug="gam-media-planner"), "T")[0]
        forge.mcps.missing_secrets = lambda entry, tok="": ["gam-key"]
        assert "do not exist yet: gam-key" in forge.name_problems(tool(secrets=["gam-key"], code=GOOD), "T")[0]
        forge.mcps.missing_secrets = lambda entry, tok="": []
        fn = forge.clean_item({"kind": "uc_function", "name": "main.tools.fx", "description": "d", "sql": "CREATE FUNCTION main.tools.fx() RETURNS INT RETURN 1"})
        calls["fn"] = "exists"
        assert "already exists" in forge.name_problems(fn, "T")[0]
        calls["fn"] = "missing"
        assert forge.name_problems(fn, "T") == []
    finally:
        forge.call, forge.mcps.act, forge.mcps.catalog, forge.mcps.missing_secrets = saved


@case("the assembled server loads in the real FastMCP, lists its tools, and refuses folders it was not given")
def _():
    py = os.environ.get("FORGE_PYTHON") or sys.executable
    probe = subprocess.run([py, "-c", "import fastmcp, openpyxl, httpx"], capture_output=True)
    if probe.returncode != 0:
        print("        (skipped: no Python with fastmcp, openpyxl and httpx; set FORGE_PYTHON)")
        return
    t = tool()
    files = forge.assemble(t, "a@x.io", "2026-10-08")
    with tempfile.TemporaryDirectory() as d:
        for rel, data in files.items():
            with open(os.path.join(d, rel), "wb") as f:
                f.write(data)
        script = r'''
import asyncio, json, sys
sys.path.insert(0, ".")
import server
from fastmcp import Client

async def main():
    async with Client(server.mcp) as c:
        tools = await c.list_tools()
        out = {t.name: (t.annotations.readOnlyHint if t.annotations else None) for t in tools}
    errors = []
    for path, allowed in (("/Volumes/main/sales/rates/x.xlsx", ("main.sales.rates",)), ("/Volumes/main/other/rates/x.xlsx", ("main.sales.rates",)),
                          ("/Volumes/main/sales/rates/../../x", ("main.sales.rates",)), ("/tmp/x", ("main.sales.rates",))):
        try:
            server._folder(path, allowed); errors.append("ok")
        except ValueError as e:
            errors.append("refused")
    try:
        server.volume_write("/Volumes/main/sales/rates/x.xlsx", b"x"); errors.append("write-ok")
    except ValueError:
        errors.append("write-refused")
    print(json.dumps({"tools": out, "paths": errors}))
asyncio.run(main())
'''
        r = subprocess.run([py, "-c", script], cwd=d, capture_output=True, text=True, timeout=120)
        assert r.returncode == 0, r.stderr[-600:]
        got = json.loads(r.stdout.strip().splitlines()[-1])
    assert got["tools"] == {"read_rate_card": True, "save_plan": False}, got
    # the read folder may be read but not written; anything outside the declared folders is refused
    assert got["paths"] == ["ok", "refused", "refused", "refused", "write-refused"], got


DRIVE_CODE = '''
import io

import openpyxl

FOLDER = "https://drive.google.com/drive/folders/1M90dZPFJjIbi51-nLotrfi_e9bDG9--i"


def save_plan(name: str) -> dict:
    """Save a one-cell plan to Drive as the next free version. name: the file name without the version."""
    try:
        taken = {f["name"] for f in drive_list("google-drive-key", FOLDER)}
        v = 1
        while "%s - v%d.xlsx" % (name, v) in taken:
            v += 1
        wb = openpyxl.Workbook()
        wb.active["A1"] = name
        buf = io.BytesIO()
        wb.save(buf)
        return {"ok": True, **drive_save("google-drive-key", FOLDER, "%s - v%d.xlsx" % (name, v), buf.getvalue())}
    except ValueError as exc:
        return {"ok": False, "error": str(exc)}
'''


def drive_tool(code=DRIVE_CODE, changes=True, secrets=("google-drive-key",)):
    return forge.clean_item({"kind": "mcp", "slug": "plan-saver", "name": "Plan saver", "description": "Saves plans to Drive.",
                             "abilities": [{"name": "save_plan", "description": "Save a plan", "changes_data": changes}],
                             "hosts": [], "secrets": list(secrets), "volumes": [], "code": code})


@case("Google Drive: the built-in helpers pass the checks, need a declared setting and a changes-data ability")
def _():
    t = drive_tool()
    assert t["problems"] == [], t["problems"]
    assert t["report"]["uses_drive"] and "www.googleapis.com" in t["report"]["hosts"], t["report"]
    files = forge.assemble(t, "a@x.io", "2026-10-09")
    server = files["server.py"].decode()
    assert "def drive_save" in server and "_DRIVE_SETTINGS = ['GOOGLE_DRIVE_KEY']" in server
    assert server.index("def drive_save") < server.index(forge.BEGIN_MARK), "the helper must sit in the fixed part"
    assert forge.model_code_of(server).strip() == DRIVE_CODE.strip()
    assert "cryptography" in files["requirements.txt"].decode()
    # a tool that does not use Drive gets none of it
    assert "drive_save" not in forge.assemble(tool(), "a@x.io", "2026-10-09")["server.py"].decode()
    assert any("must name" in p for p in drive_tool(secrets=()).get("problems")), "an undeclared setting must be refused"
    assert any("changing data" in p for p in drive_tool(changes=False)["problems"])
    bad = DRIVE_CODE + "\n\ndef drive_save(a: str) -> dict:\n    \"\"\"x\"\"\"\n    return {}\n"
    assert any("already provided" in p for p in drive_tool(code=bad)["problems"])


@case("Google Drive: the real server signs in with a service-account key, picks the next version and never overwrites")
def _():
    py = os.environ.get("FORGE_PYTHON") or sys.executable
    probe = subprocess.run([py, "-c", "import fastmcp, openpyxl, httpx, cryptography"], capture_output=True)
    if probe.returncode != 0:
        print("        (skipped: no Python with fastmcp, openpyxl, httpx and cryptography; set FORGE_PYTHON)")
        return
    files = forge.assemble(drive_tool(), "a@x.io", "2026-10-09")
    with tempfile.TemporaryDirectory() as d:
        for rel, data in files.items():
            with open(os.path.join(d, rel), "wb") as f:
                f.write(data)
        script = r'''
import base64, json, os, sys
sys.path.insert(0, ".")
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding, rsa
key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode()
os.environ["GOOGLE_DRIVE_KEY"] = json.dumps({"type": "service_account", "client_email": "bot@p.iam.gserviceaccount.com",
                                             "private_key": pem, "token_uri": "https://oauth2.googleapis.com/token"})
import server

class R:
    def __init__(self, code, body):
        self.status_code, self._b = code, body
        self.text = json.dumps(body)
    def json(self):
        return self._b

files = [{"name": "Acme - Spring - v1.xlsx", "id": "f1", "webViewLink": "https://drive/f1"}]
seen = {"uploads": [], "quota": False}
def unb64(s):
    return base64.urlsafe_b64decode(s + "=" * (-len(s) % 4))
class FakeHttpx:
    @staticmethod
    def post(url, data=None, content=None, params=None, headers=None, timeout=None):
        if url == "https://oauth2.googleapis.com/token":
            head, body, sig = data["assertion"].split(".")
            key.public_key().verify(unb64(sig), (head + "." + body).encode(), padding.PKCS1v15(), hashes.SHA256())
            claims = json.loads(unb64(body))
            assert claims["iss"] == "bot@p.iam.gserviceaccount.com" and claims["scope"].endswith("/auth/drive")
            return R(200, {"access_token": "ya29.fake", "expires_in": 3600})
        assert url.startswith("https://www.googleapis.com/upload/drive/v3/files") and headers["Authorization"] == "Bearer ya29.fake"
        assert params["supportsAllDrives"] == "true"
        meta = json.loads(content.split(b"\r\n\r\n", 2)[1].split(b"\r\n--")[0])
        if seen["quota"]:
            return R(403, {"error": {"errors": [{"reason": "storageQuotaExceeded"}]}})
        seen["uploads"].append(meta)
        files.insert(0, {"name": meta["name"], "id": "f2", "webViewLink": "https://drive/f2"})
        return R(200, {"name": meta["name"], "id": "f2", "webViewLink": "https://drive/f2"})
    @staticmethod
    def get(url, params=None, headers=None, timeout=None):
        assert "'1M90dZPFJjIbi51-nLotrfi_e9bDG9--i' in parents" in params["q"]
        if seen.get("api_off"):
            return R(403, {"error": {"errors": [{"reason": "accessNotConfigured"}]}})
        return R(200, {"files": files})
server.httpx = FakeHttpx
first = server.save_plan("Acme - Spring")
again = None
try:
    server.drive_save("google-drive-key", "1M90dZPFJjIbi51-nLotrfi_e9bDG9--i", "Acme - Spring - v1.xlsx", b"x")
except ValueError as e:
    again = str(e)
seen["quota"] = True
quota = server.save_plan("Other")
try:
    server.drive_list("some-other-key", "1M90dZPFJjIbi51-nLotrfi_e9bDG9--i"); other = "allowed"
except ValueError as e:
    other = str(e)
seen["api_off"] = True
api_off = server.save_plan("Third")
print(json.dumps({"first": first, "again": again, "quota": quota, "uploads": seen["uploads"], "other": other, "api_off": api_off}))
'''
        r = subprocess.run([py, "-c", script], cwd=d, capture_output=True, text=True, timeout=120)
        assert r.returncode == 0, r.stderr[-800:]
        got = json.loads(r.stdout.strip().splitlines()[-1])
    assert got["first"] == {"ok": True, "name": "Acme - Spring - v2.xlsx", "id": "f2", "link": "https://drive/f2"}, got
    assert got["uploads"] == [{"name": "Acme - Spring - v2.xlsx", "parents": ["1M90dZPFJjIbi51-nLotrfi_e9bDG9--i"]}], got
    assert "already in that Drive folder" in got["again"], got
    assert got["quota"]["ok"] is False and "Shared Drive" in got["quota"]["error"], got
    assert "may not use the setting" in got["other"], got
    assert "Drive API is switched off" in got["api_off"]["error"] and "project=p" in got["api_off"]["error"], got


@case("a deployment Databricks failed to copy is submitted again by itself (twice at most); a real fault is not")
def _():
    import mcps

    item = tool()
    name = forge.entry_for(item)["app_name"]
    temp = "Error downloading source code. Error: error listing files: request timed out after 1m0s of inactivity"

    def app(state, message="", ok=False):
        dep = {"deployment_id": "d1", "status": {"state": state, "message": message}}
        return {"name": name, "compute_status": {"state": "ACTIVE"}, "last_deployment_id": "d1", "active_deployment": dep,
                **({"app_status": {"state": "RUNNING"}} if ok else {})}

    def go(shown):
        seen = {"polls": 0, "redeploys": 0}

        def apps(tok):
            seen["polls"] += 1
            return {name: shown(seen)}, {}

        def redeploy(entry, tok):
            seen["redeploys"] += 1

        saved = (mcps.install, mcps._apps, mcps.redeploy, mcps.grant_volumes, forge.time.sleep)
        mcps.install, mcps._apps, mcps.redeploy = (lambda entry, who, tok, files=None: True), apps, redeploy
        mcps.grant_volumes, forge.time.sleep = (lambda entries, vols, tok: []), (lambda s: None)
        mcps._progress.pop(name, None)
        try:
            forge.install_app(item, {"user_name": "a@x.io"}, "T", "2026-10-09")
        finally:
            mcps.install, mcps._apps, mcps.redeploy, mcps.grant_volumes, forge.time.sleep = saved
        return seen, dict(mcps._progress.get(name) or {})

    # fails once with a temporary fault, then the second try works
    seen, prog = go(lambda s: app("FAILED", temp) if s["redeploys"] == 0 else app("SUCCEEDED", ok=True))
    assert seen["redeploys"] == 1 and prog["phase"] == "done", (seen, prog)
    # keeps failing the same way: two more tries, then the reason is shown as Databricks gave it
    seen, prog = go(lambda s: app("FAILED", temp))
    assert seen["redeploys"] == 2 and prog["phase"] == "failed" and "timed out" in prog["message"], (seen, prog)
    # something wrong with the tool's own files is not retried
    seen, prog = go(lambda s: app("FAILED", "requirements.txt: no matching distribution found for nosuchpkg"))
    assert seen["redeploys"] == 0 and prog["phase"] == "failed" and "nosuchpkg" in prog["message"], (seen, prog)


def main() -> int:
    passed = failed = 0
    for name, fn in CASES:
        try:
            fn()
        except Exception as exc:  # noqa: BLE001 - test harness
            print("  FAIL  " + name)
            print("        " + type(exc).__name__ + ": " + str(exc)[:300])
            failed += 1
        else:
            print("  PASS  " + name)
            passed += 1
    print("--- %d passed, %d failed ---" % (passed, failed))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
