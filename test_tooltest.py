"""Tool-testing tests: trying a new tool on real input, and repairing what breaks. No network.

Run: python test_tooltest.py

The wire shapes below are the ones seen live (2026-10-08) against deployed tools: success carries
`structuredContent`; a handled refusal is `ok: false`; a crash is `isError` with the exception text;
a call whose arguments do not fit the schema is rejected before the tool runs. The last cases start a
real generated server with the real FastMCP library (set FORGE_PYTHON to a Python that has fastmcp,
openpyxl and httpx; skipped, and says so, otherwise) and drive the whole find-the-bug, repair, retest loop.
"""
from __future__ import annotations

import json
import os
import socket
import subprocess
import sys
import tempfile
import time

import forge
import mcps
import test_designer as td
import tooltest
from dbx import DbxError

CASES = []


def case(name):
    def deco(fn):
        CASES.append((name, fn))
        return fn

    return deco


# --- a fake tool app on the wire ---------------------------------------------------------

class Resp:
    def __init__(self, body=None, status=200, headers=None, sse=True):
        self.status_code = status
        self.headers = headers or {}
        self.text = ("event: message\ndata: " + json.dumps(body) + "\n\n") if (body is not None and sse) else (json.dumps(body) if body is not None else "")


class FakeHttp:
    """Plays a tool over MCP: scripted answers per tool name, and a log of what was sent."""

    def __init__(self, answers=None, tools=None, status=200):
        self.answers, self.log, self.status = answers or {}, [], status
        self.tools = tools if tools is not None else [
            {"name": "read_sheet", "annotations": {"readOnlyHint": True}, "inputSchema": {"properties": {"name": {"type": "string"}}}},
            {"name": "save_plan", "annotations": {"readOnlyHint": False}, "inputSchema": {"properties": {"path": {"type": "string"}}}}]

    def post(self, url, headers=None, json=None, timeout=None):  # noqa: A002
        self.log.append((url, headers, json))
        if self.status != 200:
            return Resp(None, self.status)
        m, i = json["method"], json.get("id")
        if m == "initialize":
            return Resp({"jsonrpc": "2.0", "id": i, "result": {"protocolVersion": "2025-03-26"}}, headers={"mcp-session-id": "S1"})
        if m == "notifications/initialized":
            return Resp(None, 202)
        if m == "tools/list":
            return Resp({"jsonrpc": "2.0", "id": i, "result": {"tools": self.tools}})
        if m == "tools/call":
            return Resp({"jsonrpc": "2.0", "id": i, "result": self.answers[json["params"]["name"]]})
        raise AssertionError(m)


def ok(data):
    return {"isError": False, "content": [{"type": "text", "text": __import__("json").dumps(data)}], "structuredContent": data}


def crash(msg):
    return {"isError": True, "content": [{"type": "text", "text": "Error calling tool 'read_sheet': " + msg}]}


def wire(fake):
    tooltest.http = lambda: fake


TEST = {"ability": "read_sheet", "arguments": {"name": "x"}, "expect": "rows", "expect_error": False}
BAD = {**TEST, "expect_error": True}


@case("a reply is read whether it is one server-sent event or plain JSON")
def _():
    assert tooltest._data(Resp({"a": 1}))["a"] == 1
    assert tooltest._data(Resp({"a": 2}, sse=False))["a"] == 2
    assert tooltest._data(Resp(None)) == {} and tooltest._data(Resp("not json", sse=False)) == {}


@case("the client opens a session, lists abilities and calls one, carrying the session and the admin's token")
def _():
    f = FakeHttp({"read_sheet": ok({"ok": True, "rows": [1]})})
    wire(f)
    c = tooltest.Client("https://tool.example/", "ADMIN").open()
    assert [t["name"] for t in c.tools()] == ["read_sheet", "save_plan"]
    r = c.call("read_sheet", {"name": "x"})
    assert r["data"] == {"ok": True, "rows": [1]} and not r["is_error"]
    assert all(h["Authorization"] == "Bearer ADMIN" for _, h, _ in f.log)
    assert f.log[-1][0] == "https://tool.example/mcp" and f.log[-1][1]["mcp-session-id"] == "S1"
    assert [b["method"] for _, _, b in f.log] == ["initialize", "notifications/initialized", "tools/list", "tools/call"]


@case("a tool that cannot be reached says so plainly, and is not mistaken for a bug in it")
def _():
    for status, fragment in ((403, "not allowed"), (401, "not allowed"), (502, "did not answer"), (404, "did not answer")):
        wire(FakeHttp(status=status))
        try:
            tooltest.Client("https://t", "T").open()
        except DbxError as exc:
            assert fragment in str(exc), (status, str(exc))
        else:
            raise AssertionError(status)


@case("connect refuses a tool that is not running yet")
def _():
    tooltest.call = lambda m, p, tok, **k: {"url": "https://t", "description": "x"}
    mcps._state = lambda app, prog: ("deploying", "Installing it.")
    try:
        tooltest.connect("rate-reader", "T")
    except DbxError as exc:
        assert exc.status == 409 and "not running" in str(exc)
    else:
        raise AssertionError("must refuse")
    mcps._state = lambda app, prog: ("running", "")
    wire(FakeHttp())
    assert tooltest.connect("rate-reader", "T").sid == "S1"


def verdict(res, test=TEST):
    return tooltest.judge(test, {"is_error": False, "text": "", "data": None, "rpc_error": None, **res}, "T")


@case("judging by rule: a crash fails, a wrongly refused call fails, a refusal of bad input passes")
def _():
    v = verdict({"is_error": True, "text": "Error calling tool 'read_sheet': 'MergedCell' object has no attribute 'column_letter'"})
    assert v["verdict"] == "fail" and v["kind"] == "crash" and v["reason"] == "It crashed: 'MergedCell' object has no attribute 'column_letter'"
    v = verdict({"data": {"ok": False, "error": "File not found"}, "text": '{"ok": false}'})
    assert v["verdict"] == "fail" and v["kind"] == "refused" and "File not found" in v["reason"]
    v = verdict({"data": {"ok": False, "error": "No such sheet"}}, BAD)
    assert v["verdict"] == "pass" and v["kind"] == "refused"
    assert verdict({"is_error": True, "text": "Error calling tool 'x': boom"}, BAD)["verdict"] == "fail"   # bad input must never crash it


@case("a test whose input does not fit is not a verdict on the tool")
def _():
    v = verdict({"is_error": True, "text": "5 validation errors for call[read_sheet]\nname\n  Input should be a valid string"})
    assert v["verdict"] == "unchecked" and v["kind"] == "input" and "says nothing about the tool" in v["reason"]
    assert verdict({"rpc_error": {"code": -32602, "message": "bad"}})["verdict"] == "unchecked"


@case("a working answer is only questioned if it holds nothing; the model's opinion is not used")
def _():
    v = verdict({"data": {"ok": True, "sheets": [{"line_items": [{"a": 1}]}]}})
    assert v["verdict"] == "pass" and v["reason"] == "It answered with 1 sheet."
    assert verdict({"data": {"ok": True, "sheets": [1, 2, 3, 4, 5], "line_items": [1] * 2}})["reason"] == "It answered with 5 sheets."
    assert verdict({"data": {"ok": True, "totals": {"total": 5}}})["reason"] == "It answered with a result."
    assert tooltest.plain_answer([1, 2]) == "2 items" and tooltest.plain_answer({"rate_cards": ["a"]}) == "1 rate card"
    v = verdict({"data": {"ok": True, "sheets": [{"line_items": [], "problems": []}], "truncated": False}})
    assert v["verdict"] == "look" and v["kind"] == "empty"
    assert verdict({"data": {"ok": True, "total": 0}})["verdict"] == "pass"           # a zero is something
    v = verdict({"data": {"ok": True}}, BAD)
    assert v["verdict"] == "look" and v["kind"] == "accepted"                            # odd, but not worth rewriting code over
    assert tooltest.substance({"ok": True, "note": "x", "warnings": ["w"], "count": 3}) == 0
    assert tooltest.substance(["a", "", None, {"b": 2}]) == 2


@case("the shape of an answer is summarised without cutting it off mid-word")
def _():
    s = tooltest.summarize({"ok": True, "sheets": [{"sheet": "A", "line_items": [{"r": i} for i in range(14)], "problems": []}], "name": "x" * 300})
    assert "14 items" in s and "[] (empty)" in s and "..." in s and len(s) < 500
    assert tooltest.summarize([]) == "[] (empty)"


TOOLS = FakeHttp().tools
ITEM = {"volumes": [{"volume": "main.sales.rates", "access": "read"}]}


@case("only tests that can really be run are kept; what changes things is listed as not run, with the reason")
def _():
    raw = {"tests": [
        {"ability": "read_sheet", "arguments": {"name": "Rates"}, "expect": "rows"},
        {"ability": "read_sheet", "arguments": {"name": "nope"}, "expect": "refusal", "expect_error": True},
        {"ability": "save_plan", "arguments": {"path": "/Volumes/main/sales/rates/x.xlsx"}},          # changes things: never run
        {"ability": "invented", "arguments": {}},                                                      # no such ability
        {"ability": "read_sheet", "arguments": {"sheet": "Rates"}},                                    # not an argument it has
        {"ability": "read_sheet", "arguments": {"name": "/Volumes/main/other/secret/x.xlsx"}},         # outside its folders
        {"ability": "read_sheet", "arguments": {"name": "/Volumes/main/sales/rates/../../x"}},        # tries to climb out
        "junk"]}
    out = tooltest.clean_plan(raw, TOOLS, ITEM)
    assert [(t["ability"], t["arguments"]["name"], t["expect_error"]) for t in out["tests"]] == [("read_sheet", "Rates", False), ("read_sheet", "nope", True)]
    assert out["skipped"] == [{"ability": "save_plan", "why": "It saves or changes something, so it is never run without you."}]
    ok_path = {"tests": [{"ability": "read_sheet", "arguments": {"name": "/Volumes/main/sales/rates/a.xlsx"}}]}
    assert len(tooltest.clean_plan(ok_path, TOOLS, ITEM)["tests"]) == 1
    # a read-only ability that got no test is reported, not silently dropped
    out = tooltest.clean_plan({"tests": [], "skipped": [{"ability": "read_sheet", "why": "no file to try it on"}]}, TOOLS, ITEM)
    assert {"ability": "read_sheet", "why": "no file to try it on"} in out["skipped"]
    out = tooltest.clean_plan({"tests": []}, TOOLS, ITEM)
    assert any(s["ability"] == "read_sheet" and "No realistic call" in s["why"] for s in out["skipped"])
    many = {"tests": [{"ability": "read_sheet", "arguments": {"name": str(i)}} for i in range(30)]}
    assert len(tooltest.clean_plan(many, TOOLS, ITEM)["tests"]) == tooltest.MAX_TESTS


@case("the folders the tool declared are listed, one level down, as the admin")
def _():
    seen = []

    def call(method, path, tok, **kw):
        seen.append((path, tok))
        listing = {"/api/2.0/fs/directories/Volumes/main/sales/rates": [{"name": "a.xlsx"}, {"name": "old", "is_directory": True}],
                   "/api/2.0/fs/directories/Volumes/main/sales/rates/old": [{"name": "b.xlsx"}]}
        if path not in listing:
            raise DbxError("nope", 404)
        return {"contents": listing[path]}

    tooltest.call = call
    f = tooltest.facts_for({"volumes": [{"volume": "main.sales.rates", "access": "read"}, {"volume": "main.sales.gone", "access": "read"}]}, "ADMIN")
    assert f["/Volumes/main/sales/rates"] == {"files": ["a.xlsx"], "folders": ["old/"], "inside": {"old/": ["b.xlsx"]}}
    assert f["/Volumes/main/sales/gone"]["files"] == []   # not readable: empty, not an error
    assert all(t == "ADMIN" for _, t in seen)


@case("the plan is written by the model from real facts, then cleaned")
def _():
    plan = {"tests": [{"ability": "read_sheet", "arguments": {"name": "Rates"}, "expect": "rows"}], "skipped": []}
    m = td.Model(td.reply(text=json.dumps(plan)))
    td.install(m)
    out = tooltest.plan(dict(ITEM, name="Rates", description="d"), TOOLS, {"/Volumes/main/sales/rates": {"files": ["a.xlsx"]}}, "T")
    assert out["tests"][0]["arguments"] == {"name": "Rates"} and out["facts"]["/Volumes/main/sales/rates"]["files"] == ["a.xlsx"]
    sent = json.loads(m.sent[0][1][1]["content"])
    assert sent["facts"]["/Volumes/main/sales/rates"]["files"] == ["a.xlsx"] and sent["abilities"][1]["read_only"] is False
    assert "never invent a file" in m.sent[0][1][0]["content"] and "SHAPE of a good answer" in m.sent[0][1][0]["content"]


# --- repairing ---------------------------------------------------------------------------------

BUGGY = '''
import io

import openpyxl


def read_sheet(file_path: str) -> dict:
    """Read the first sheet of a workbook. file_path: the file to read."""
    ws = openpyxl.load_workbook(io.BytesIO(volume_read(file_path))).active
    return {"ok": True, "columns": [c.column_letter for c in ws[1]]}
'''
FIXED = BUGGY.replace("import openpyxl\n", "import openpyxl\nfrom openpyxl.utils import get_column_letter\n").replace("c.column_letter", "get_column_letter(c.column)")
FITEM = forge.clean_item({"kind": "mcp", "slug": "sheet-reader", "name": "Sheet reader", "description": "Reads sheets.",
                          "abilities": [{"name": "read_sheet", "description": "Reads.", "changes_data": False}], "code": FIXED,
                          "volumes": [{"volume": "main.sales.rates", "access": "read"}]})
BROKEN = forge.clean_item({**FITEM, "code": BUGGY})
FAIL = [{"ability": "read_sheet", "arguments": {"file_path": "/Volumes/main/sales/rates/a.xlsx"},
         "reason": "It crashed: 'MergedCell' object has no attribute 'column_letter'"}]


def python_block(code, what="Merged cells have no column letter, so the code now asks for it another way."):
    return td.reply(text="```python\n" + code + "```\nWHAT CHANGED: " + what)


@case("a repair is proposed, must pass the same checks, and comes back as a diff with its own fingerprint")
def _():
    m = td.Model(python_block(FIXED))
    td.install(m)
    out = tooltest.propose_fix(BROKEN, FAIL, "T")
    assert out["problems"] == [] and out["code"].strip() == FIXED.strip() and not out["no_change"]
    assert out["fingerprint"] != BROKEN["fingerprint"] and out["fingerprint"] == forge.clean_item({**BROKEN, "code": FIXED})["fingerprint"]
    assert out["what"].startswith("Merged cells have no column letter")
    assert any(l.startswith("-") and "column_letter" in l for l in out["diff"]) and any(l.startswith("+") and "get_column_letter" in l for l in out["diff"])
    sent = json.loads(m.sent[0][1][1]["content"])
    assert sent["current_code"].strip() == BUGGY.strip() and "MergedCell" in sent["failures"][0]["what_happened"]
    assert "repairing a tool" in m.sent[0][1][0]["content"] and "keep every ability's name" in m.sent[0][1][0]["content"]


@case("a repair that breaks the rules is sent back with the reasons, and one that changes nothing is refused")
def _():
    still_bad = BUGGY + "\n\ndef extra(x: int) -> int:\n    \"\"\"x\"\"\"\n    return x.column_letter\n"
    m = td.Model(python_block(still_bad), python_block(FIXED))
    td.install(m)
    out = tooltest.propose_fix(BROKEN, FAIL, "T")
    assert out["problems"] == [] and out["code"].strip() == FIXED.strip()
    retry = json.loads(m.sent[1][1][1]["content"])
    assert any("fails on merged cells" in p for p in retry["fix_these_first"])
    m = td.Model(python_block(LIVE), python_block(LIVE), python_block(LIVE))
    td.install(m)
    out = tooltest.propose_fix(LIVE_ITEM, FAIL, "T")
    assert out["problems"] == ["The code is the same as before: change what is wrong."]
    # and code the checks already refuse is reported with their reason, never passed on
    m = td.Model(python_block(BUGGY), python_block(BUGGY), python_block(BUGGY))
    td.install(m)
    assert "fails on merged cells" in " ".join(tooltest.propose_fix(BROKEN, FAIL, "T")["problems"])


@case("the repair model may say the code is fine and the test was at fault")
def _():
    m = td.Model(td.reply(text="NO CHANGE NEEDED: the test asked for a file that is not in the folder."))
    td.install(m)
    out = tooltest.propose_fix(BROKEN, FAIL, "T")
    assert out["no_change"] and out["diff"] == [] and out["problems"] == [] and "not in the folder" in out["what"]
    assert out["code"] == BROKEN["code"]


# --- reading a tool back, and the routes -----------------------------------------------------------

@case("a tool is rebuilt from the files it was deployed with, identical to how it was first described")
def _():
    files = forge.assemble(FITEM, "a@x.io", "2026-10-08")
    back = forge.item_from_files(files, "sheet-reader")
    assert back["code"].strip() == FITEM["code"].strip() and back["fingerprint"] == FITEM["fingerprint"]
    assert [a["name"] for a in back["abilities"]] == ["read_sheet"] and back["volumes"] == FITEM["volumes"]
    assert forge.model_code_of("print('not one of ours')") == "" and forge.item_from_files({"server.py": b"x"}, "s") is None
    assert forge.item_from_files({"mcp.yaml": b"name: x", "server.py": b"print(1)"}, "s") is None


@case("only a tool this designer made can be read back for testing or repair")
def _():
    files = forge.assemble(FITEM, "a@x.io", "2026-10-08")
    forge.mcps._export = lambda path, tok: files[path.rsplit("/", 1)[1]]
    forge.call = lambda m, p, tok, **k: {"description": forge.GEN_MARK + " Reads sheets."}
    assert forge.item_from_workspace("sheet-reader", "T")["fingerprint"] == FITEM["fingerprint"]
    forge.call = lambda m, p, tok, **k: {"description": "Someone's own tool"}
    for slug, status in (("sheet-reader", 403), ("../etc", 400)):
        try:
            forge.item_from_workspace(slug, "T")
        except DbxError as exc:
            assert exc.status == status, (slug, exc.status)
        else:
            raise AssertionError(slug)


def routes(item=None):
    portal, c = td.client(admin=True)
    files = forge.assemble(item or FITEM, "a@x.io", "2026-10-08")
    forge.mcps._export = lambda path, tok: files[path.rsplit("/", 1)[1]]
    forge.call = lambda m, p, tok, **k: {"description": forge.GEN_MARK + " Reads sheets."}
    return portal, c


@case("every tool-test route refuses someone who is not an admin")
def _():
    _, c = td.client(admin=False)
    for path in ("plan", "run", "fix", "apply"):
        r = c.post("/api/admin/designer/tool-test/" + path, headers=td.H, json={"slug": "sheet-reader"})
        assert r.status_code == 403, (path, r.status_code)


@case("a test run only calls a read-only ability, takes its read-only status from the tool, and rejects what is not its own")
def _():
    portal, c = routes()
    f = FakeHttp({"read_sheet": ok({"ok": True, "rows": ["x"]})})
    wire(f)
    tooltest.call = lambda m, p, tok, **k: {"url": "https://t", "description": "x"}
    mcps._state = lambda app, prog: ("running", "")
    body = {"slug": "sheet-reader", "test": {"ability": "read_sheet", "arguments": {"name": "x"}, "expect": "rows"}}
    r = c.post("/api/admin/designer/tool-test/run", headers=td.H, json=body)
    assert r.status_code == 200 and r.json()["verdict"] == "pass", r.text
    # the request says the ability is harmless; the tool's own declaration decides
    r = c.post("/api/admin/designer/tool-test/run", headers=td.H, json={"slug": "sheet-reader", "test": {"ability": "save_plan", "arguments": {}, "readonly": True}})
    assert r.status_code == 400 and "saves or changes" in r.json()["detail"]
    assert c.post("/api/admin/designer/tool-test/run", headers=td.H, json={"slug": "sheet-reader", "test": {"ability": "nope"}}).status_code == 400
    assert c.post("/api/admin/designer/tool-test/run", headers=td.H, json={"slug": "Bad Slug", "test": {}}).status_code == 400
    assert [b["params"]["name"] for _, _, b in f.log if b["method"] == "tools/call"] == ["read_sheet"]   # save_plan was never called


@case("the fix route proposes and deploys nothing; apply needs the exact fingerprint, checks the code again, and deploys in the background")
def _():
    portal, c = routes(BROKEN)   # the buggy tool is what is deployed
    td.install(td.Model(python_block(FIXED)))
    fixes = c.post("/api/admin/designer/tool-test/fix", headers=td.H, json={"slug": "sheet-reader", "failures": FAIL})
    assert fixes.status_code == 200 and fixes.json()["code"].strip() == FIXED.strip() and fixes.json()["diff"]
    assert c.post("/api/admin/designer/tool-test/fix", headers=td.H, json={"slug": "sheet-reader", "failures": []}).status_code == 400
    deployed = []
    forge.install_app = lambda item, who, tok, date: deployed.append(item["code"])
    fp = fixes.json()["fingerprint"]
    body = {"slug": "sheet-reader", "code": FIXED, "approved": fp, "what": "x"}
    assert c.post("/api/admin/designer/tool-test/apply", headers=td.H, json={**body, "approved": "0" * 64}).status_code == 400
    assert c.post("/api/admin/designer/tool-test/apply", headers=td.H, json={**body, "code": FIXED + "\n# edited after approval\n"}).status_code == 400
    assert c.post("/api/admin/designer/tool-test/apply", headers=td.H, json={**body, "code": BUGGY}).status_code == 400   # nothing to change
    evil = BUGGY + "\n\ndef boom(x: int) -> int:\n    \"\"\"x\"\"\"\n    return eval('1')\n"
    assert c.post("/api/admin/designer/tool-test/apply", headers=td.H, json={**body, "code": evil, "approved": forge.clean_item({**BROKEN, "code": evil})["fingerprint"]}).status_code == 400
    assert deployed == []                                # nothing got through
    r = c.post("/api/admin/designer/tool-test/apply", headers=td.H, json=body)
    assert r.status_code == 200 and r.json()["started"] is True and r.json()["app_name"] == "mcp-sheet-reader"
    assert r.json()["since"] <= time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())     # a time to wait for a newer deployment than
    assert deployed == [FIXED] or [d.strip() for d in deployed] == [FIXED.strip()]


@case("a repair is only 'live' once nothing is pending and the active deployment is newer than the moment it was applied")
def _():
    since = "2026-10-08T12:00:00Z"
    app = {"app_status": {"state": "RUNNING"}, "active_deployment": {"status": {"state": "SUCCEEDED"}, "create_time": "2026-10-08T11:00:00Z"}}
    tooltest.call = lambda m, p, tok, **k: app
    mcps._progress.clear()
    old = tooltest.deployed_since("sheet-reader", since, "T")
    assert old["ready"] is False and old["failed"] is False and "Installing" in old["note"]    # still the OLD version, though RUNNING
    app["pending_deployment"] = {"status": {"state": "IN_PROGRESS"}}
    app["active_deployment"] = {"status": {"state": "SUCCEEDED"}, "create_time": "2026-10-08T12:00:05Z"}
    assert tooltest.deployed_since("sheet-reader", since, "T")["ready"] is False                # new one landed but another is pending
    del app["pending_deployment"]
    assert tooltest.deployed_since("sheet-reader", since, "T") == {"ready": True, "failed": False, "note": ""}
    app["pending_deployment"] = {"status": {"state": "FAILED", "message": "build failed"}}
    out = tooltest.deployed_since("sheet-reader", since, "T")
    assert out["failed"] and "build failed" in out["note"]
    del app["pending_deployment"]
    mcps._progress["mcp-sheet-reader"] = {"phase": "failed", "message": "Could not copy", "at": time.time()}
    assert tooltest.deployed_since("sheet-reader", since, "T")["failed"]
    mcps._progress.clear()
    mcps._progress["mcp-sheet-reader"] = {"phase": "copying", "message": "x", "at": time.time()}
    app["active_deployment"]["create_time"] = "2026-10-08T11:00:00Z"
    assert "Copying" in tooltest.deployed_since("sheet-reader", since, "T")["note"]
    mcps._progress.clear()


@case("the ready route is admin-only, checks its inputs, and answers for the tool it names")
def _():
    _, c = td.client(admin=False)
    assert c.get("/api/admin/designer/tool-test/ready?slug=a&since=2026-10-08T12:00:00Z", headers=td.H).status_code == 403
    _, c = td.client(admin=True)
    assert c.get("/api/admin/designer/tool-test/ready?slug=Bad Slug&since=2026-10-08T12:00:00Z", headers=td.H).status_code == 400
    assert c.get("/api/admin/designer/tool-test/ready?slug=sheet-reader&since=yesterday", headers=td.H).status_code == 400
    tooltest.call = lambda m, p, tok, **k: {"app_status": {"state": "RUNNING"}, "active_deployment": {"status": {"state": "SUCCEEDED"}, "create_time": "2026-10-08T12:01:00Z"}}
    r = c.get("/api/admin/designer/tool-test/ready?slug=sheet-reader&since=2026-10-08T12:00:00Z", headers=td.H)
    assert r.status_code == 200 and r.json()["ready"] is True


# --- the real thing: a generated server, the real FastMCP, the whole loop ----------------------------------

def free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class Running:
    """A generated server run as a real process, so the client talks real MCP to it."""

    def __init__(self, py, item):
        self.dir = tempfile.TemporaryDirectory()
        for rel, data in forge.assemble(item, "a@x.io", "2026-10-08").items():
            with open(os.path.join(self.dir.name, rel), "wb") as f:
                f.write(data)
        self.port = free_port()
        self.proc = subprocess.Popen([py, "server.py"], cwd=self.dir.name, env={**os.environ, "PORT": str(self.port)},
                                     stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(80):
            try:
                with socket.create_connection(("127.0.0.1", self.port), timeout=0.5):
                    break
            except OSError:
                time.sleep(0.25)
        else:
            raise AssertionError("the generated server did not start")

    def stop(self):
        self.proc.terminate()
        self.proc.wait(timeout=10)
        self.dir.cleanup()


class Direct:
    """The real HTTP client, pointed at a local server."""

    @staticmethod
    def post(url, headers=None, json=None, timeout=None):  # noqa: A002
        import httpx

        r = httpx.post(url, headers=headers, json=json, timeout=60)
        return r


LIVE = '''
def read_sheet(name: str = "") -> dict:
    """Read a sheet by name. name: the sheet to read."""
    if name == "missing":
        return {"ok": False, "error": "No sheet called missing."}
    return {"ok": True, "sheet": name or "first", "rows": [["a", 1]]}


def read_cells(name: str = "") -> dict:
    """Read the cells of a sheet. name: the sheet to read."""
    cells = [{"col": "A"}, {"col": "B"}]
    return {"ok": True, "letters": [c["column"] for c in cells]}
'''
LIVE_FIXED = LIVE.replace('c["column"]', 'c["col"]')
LIVE_ITEM = forge.clean_item({"kind": "mcp", "slug": "live-sheet", "name": "Live sheet", "description": "Reads sheets.",
                              "abilities": [{"name": "read_sheet", "description": "Reads.", "changes_data": False},
                                            {"name": "read_cells", "description": "Cells.", "changes_data": False}], "code": LIVE})


@case("against a real server: the bug is found, the working ability is not blamed, then the repair is proved by running again")
def _():
    py = os.environ.get("FORGE_PYTHON") or sys.executable
    if subprocess.run([py, "-c", "import fastmcp, httpx"], capture_output=True).returncode != 0:
        print("        (skipped: no Python with fastmcp and httpx; set FORGE_PYTHON)")
        return
    assert LIVE_ITEM["problems"] == [], LIVE_ITEM["problems"]
    tooltest.http = lambda: Direct
    tests = [{"ability": "read_sheet", "arguments": {"name": "Rates"}, "expect": "rows", "expect_error": False},
             {"ability": "read_sheet", "arguments": {"name": "missing"}, "expect": "refusal", "expect_error": True},
             {"ability": "read_cells", "arguments": {"name": "Rates"}, "expect": "letters", "expect_error": False}]
    srv = Running(py, LIVE_ITEM)
    try:
        c = tooltest.Client("http://127.0.0.1:%d" % srv.port, "").open()
        by = {t["name"]: t for t in c.tools()}
        assert set(by) == {"read_sheet", "read_cells"} and all(tooltest._readonly(t) for t in by.values())
        results = [tooltest.run_one(c, t, "T") for t in tests]
    finally:
        srv.stop()
    assert [(r["ability"], r["verdict"]) for r in results] == [("read_sheet", "pass"), ("read_sheet", "pass"), ("read_cells", "fail")]
    crash_ = results[2]
    assert crash_["kind"] == "crash" and "column" in crash_["reason"]
    # the repair is proposed from exactly what the test saw, shown as a diff, and passes the same checks
    td.install(td.Model(python_block(LIVE_FIXED, "It asked for a key the data does not have; it now asks for the right one.")))
    fix = tooltest.propose_fix(LIVE_ITEM, [crash_], "T")
    assert fix["problems"] == [] and any('c["col"]' in l for l in fix["diff"] if l.startswith("+"))
    # deploy it (here: start it again) and run the same tests again
    repaired = forge.clean_item({**LIVE_ITEM, "code": fix["code"]})
    srv = Running(py, repaired)
    try:
        c = tooltest.Client("http://127.0.0.1:%d" % srv.port, "").open()
        again = [tooltest.run_one(c, t, "T") for t in tests]
    finally:
        srv.stop()
    assert [r["verdict"] for r in again] == ["pass", "pass", "pass"], again


@case("on a deployed portal the tool is checked even though the user token has no apps scope, and called as the portal if it refuses the user")
def _():
    import os as _os
    calls, posts = [], []

    def call(method, path, tok, **kw):
        calls.append((method, path, tok))
        if path == "/api/2.0/apps/mcp-sheet-reader" and tok == "USER":
            raise DbxError("Provided OAuth token does not have required scopes: apps", 403)
        if path == "/api/2.0/apps/mcp-sheet-reader":
            return {"name": "mcp-sheet-reader", "url": "https://tool.apps", "compute_status": {"state": "ACTIVE"},
                    "app_status": {"state": "RUNNING"}, "active_deployment": {"status": {"state": "SUCCEEDED"}}}
        if path.startswith("/api/2.0/permissions/apps/"):
            return {}
        raise AssertionError(path)

    class Resp:
        def __init__(self, code):
            self.status_code, self.headers, self.text = code, {"mcp-session-id": "s1"}, "{}"

        def json(self):
            return {"jsonrpc": "2.0", "id": 1, "result": {}}

    class Http:
        def post(self, url, headers=None, json=None, timeout=None):
            posts.append(headers.get("Authorization"))
            return Resp(403 if headers.get("Authorization") == "Bearer USER" else 200)

    saved = (tooltest.call, tooltest.http, tooltest.in_apps, tooltest.app_token, _os.environ.get("DATABRICKS_CLIENT_ID"))
    tooltest.call, tooltest.http, tooltest.in_apps, tooltest.app_token = call, lambda: Http(), lambda: True, lambda: "APP"
    _os.environ["DATABRICKS_CLIENT_ID"] = "portal-sp"
    try:
        client = tooltest.connect("sheet-reader", "USER")
    finally:
        tooltest.call, tooltest.http, tooltest.in_apps, tooltest.app_token = saved[:4]
        if saved[4] is None:
            _os.environ.pop("DATABRICKS_CLIENT_ID", None)
        else:
            _os.environ["DATABRICKS_CLIENT_ID"] = saved[4]
    assert ("GET", "/api/2.0/apps/mcp-sheet-reader", "APP") in calls          # read as the portal after the scope refusal
    assert ("PATCH", "/api/2.0/permissions/apps/mcp-sheet-reader", "USER") in calls  # the portal is given Can use, as the admin
    assert posts[0] == "Bearer USER" and posts[-1] == "Bearer APP" and client.tok == "APP"


def main() -> int:
    passed = failed = 0
    for name, fn in CASES:
        try:
            fn()
        except Exception as exc:  # noqa: BLE001 - test harness
            print("  FAIL  " + name)
            print("        " + type(exc).__name__ + ": " + str(exc)[:300])
            failed += 1
        else:
            print("  PASS  " + name)
            passed += 1
    print("--- %d passed, %d failed ---" % (passed, failed))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
